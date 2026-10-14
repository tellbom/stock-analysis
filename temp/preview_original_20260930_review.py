"""Read-only grid preview when the artifact renderer is unavailable."""
from pathlib import Path
from datetime import datetime, date
import math
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'outputs/01a116ce-1f0c-7d41-8967-adcca880c1b8/2026-09-30_原表纠正回测_20261014/2026-09-30原名单_10月8日入场_纠正回测.xlsx'
comparison = '--comparison' in sys.argv
weekly = '--weekly' in sys.argv
qfq_weekly = '--qfq-weekly' in sys.argv
if qfq_weekly:
    weekly = True
    sys.argv.remove('--qfq-weekly')
if weekly:
    SOURCE = ROOT / 'outputs/01a116ce-1f0c-7d41-8967-adcca880c1b8/2026-10-14_统一基线_D3_D5/2026-10-14_D3_D5_普通与混合选股对照.xlsx'
    if '--weekly' in sys.argv:
        sys.argv.remove('--weekly')
    if qfq_weekly:
        SOURCE = ROOT / 'outputs/01a116ce-1f0c-7d41-8967-adcca880c1b8/2026-10-14_复权修复_D3_D5/2026-10-14_D3_D5_复权修复_普通与混合.xlsx'
if comparison:
    SOURCE = ROOT / 'outputs/01a116ce-1f0c-7d41-8967-adcca880c1b8/原版统一基线对照_20260930/普通与混合D3_D5_统一基线对照.xlsx'
    sys.argv.remove('--comparison')
book = load_workbook(SOURCE, data_only=True)
font_root = Path('C:/Windows/Fonts')
clipped = []


def rgb(color, fallback):
    return '#' + color.rgb[-6:] if color is not None and color.type == 'rgb' else fallback


def display(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float):
        if '%' in cell.number_format:
            prefix = cell.number_format.split('%')[0]
            digits = len(prefix.split('.')[-1]) if '.' in prefix else 0
            return f'{value:.{digits}%}'
        if cell.number_format == '0.00':
            return f'{value:.2f}'
        if cell.number_format == '0.000000':
            return f'{value:.6f}'
        if cell.number_format == '#,##0.00':
            return f'{value:,.2f}'
        if value.is_integer():
            return str(int(value))
    return str(value)


views = [('对照结果', 9, 35), ('训练口径', 7, 58), ('逐股收益', 12, 16), ('行情输入', 8, 16)] if comparison else [('原表回测', 10, 35), ('逐股收益', 11, 16), ('行情输入', 8, 18)]
if weekly:
    views = [('本轮口径', 8, 39), ('D3选股对照', 13, 18), ('D5选股对照', 13, 18), ('Gate诊断', 9, 16), ('训练因子', 7, 55), ('行情与覆盖', 11, 16)]
    if qfq_weekly:
        views = [('本轮口径', 8, 39), ('D3选股对照', 15, 18), ('D5选股对照', 15, 18), ('Gate诊断', 9, 16), ('训练因子', 7, 55), ('行情与覆盖', 11, 16)]
for name, cols, rows in views:
    if len(sys.argv) > 1 and name not in sys.argv[1:]:
        continue
    sheet = book[name]
    widths = []
    for c in range(1, cols + 1):
        dim = next((d for d in sheet.column_dimensions.values() if d.min <= c <= d.max), None)
        widths.append(int((dim.width if dim else 13) * 7 + 5))
    heights = [int((sheet.row_dimensions[r].height or 15) * 4 / 3) for r in range(1, rows + 1)]
    xs, ys = [14], [12]
    for width in widths:
        xs.append(xs[-1] + width)
    for height in heights:
        ys.append(ys[-1] + height)
    image = Image.new('RGB', (xs[-1] + 14, ys[-1] + 12), 'white')
    draw = ImageDraw.Draw(image)
    cells = []
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            cell = sheet.cell(r, c)
            merged = next((m for m in sheet.merged_cells.ranges if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col), None)
            if merged and (r != merged.min_row or c != merged.min_col):
                continue
            last_col = min(merged.max_col, cols) if merged else c
            last_row = min(merged.max_row, rows) if merged else r
            box = (xs[c-1], ys[r-1], xs[last_col], ys[last_row])
            if cell.fill.patternType == 'solid':
                draw.rectangle(box, fill=rgb(cell.fill.fgColor, '#FFFFFF'))
            if cell.border.bottom.style:
                draw.line((box[0], box[3], box[2], box[3]), fill=rgb(cell.border.bottom.color, '#A9BBCD'), width=1)
            if cell.value is not None:
                cells.append((cell, box, c, r, bool(merged)))
    for cell, box, c, r, merged in cells:
        text = display(cell)
        size = round((cell.font.sz or 10) * 4 / 3)
        font = ImageFont.truetype(str(font_root / ('msyhbd.ttc' if cell.font.bold else 'msyh.ttc')), size)
        color = rgb(cell.font.color, '#24364B')
        available = box[2] - box[0] - 8
        if not cell.alignment.wrap_text and not merged and not isinstance(cell.value, (int, float)):
            next_col = c + 1
            while next_col <= cols and sheet.cell(r, next_col).value is None:
                available += widths[next_col-1]
                next_col += 1
        lines = [text]
        if cell.alignment.wrap_text:
            lines, current = [], ''
            for char in text:
                if current and draw.textlength(current + char, font=font) > available:
                    lines.append(current)
                    current = ''
                current += char
            lines.append(current)
        elif draw.textlength(text, font=font) > available + 1:
            clipped.append(f'{name}!{cell.coordinate}: horizontal')
        line_height = size + 3
        if len(lines) * line_height > box[3] - box[1] - 2:
            clipped.append(f'{name}!{cell.coordinate}: vertical')
        y = box[1] + max(2, (box[3] - box[1] - len(lines) * line_height) / 2)
        for line in lines:
            x = box[0] + 4
            if cell.alignment.horizontal == 'center':
                x = (box[0] + box[2] - draw.textlength(line, font=font)) / 2
            elif cell.alignment.horizontal == 'right' or isinstance(cell.value, (int, float)):
                x = box[2] - draw.textlength(line, font=font) - 4
            draw.text((x, y), line, font=font, fill=color)
            y += line_height
    output = ROOT / 'temp' / f'{"weekly" if weekly else "controlled" if comparison else "original"}_grid_{name}.png'
    image.save(output)
    print(output)
print('Potential clipping:', clipped)
